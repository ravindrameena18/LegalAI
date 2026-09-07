import io
import logging
from dataclasses import dataclass, field

import docx
import pypdf

logger = logging.getLogger("legalai.extractor")


@dataclass
class ExtractedPage:
    page_number: int
    text: str


@dataclass
class ExtractionResult:
    success: bool
    status: str  # "ready", "ocr_required", "failed", "password_required"
    page_count: int
    pages: list[ExtractedPage] = field(default_factory=list)
    error_message: str | None = None
    is_encrypted: bool = False
    unencrypted_content: bytes | None = None


def extract_document_text(content: bytes, file_type: str, password: str | None = None) -> ExtractionResult:
    """
    Extract structured text from supported documents (PDF, DOCX, TXT).
    Detects password-protected PDFs, scanned/image-only PDFs (flags as ocr_required),
    and decrypts password-protected PDFs when valid password is provided.
    """
    file_type = file_type.lower().strip()

    try:
        # 1. Plain Text extraction
        if file_type == "txt":
            try:
                text = content.decode("utf-8")
            except UnicodeDecodeError:
                text = content.decode("latin-1")
            return ExtractionResult(
                success=True,
                status="ready",
                page_count=1,
                pages=[ExtractedPage(page_number=1, text=text)],
            )

        # 2. DOCX extraction
        if file_type == "docx":
            doc = docx.Document(io.BytesIO(content))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)

            combined_text = "\n\n".join(paragraphs)
            return ExtractionResult(
                success=True,
                status="ready",
                page_count=1,
                pages=[ExtractedPage(page_number=1, text=combined_text)],
            )

        # 3. PDF extraction
        if file_type == "pdf":
            reader = pypdf.PdfReader(io.BytesIO(content))
            is_encrypted = getattr(reader, "is_encrypted", False)
            unencrypted_content: bytes | None = None

            if is_encrypted:
                decrypted = False
                if password:
                    try:
                        res = reader.decrypt(password)
                        if res != 0:
                            decrypted = True
                    except Exception as decrypt_exc:
                        logger.warning("Error decrypting PDF with password: %s", decrypt_exc)
                else:
                    # Attempt empty password decryption (e.g. permissions-only protected)
                    try:
                        res = reader.decrypt("")
                        if res != 0:
                            decrypted = True
                    except Exception:
                        pass

                if not decrypted:
                    logger.info("PDF is password-protected and requires password to unlock.")
                    return ExtractionResult(
                        success=True,
                        status="password_required",
                        page_count=0,
                        pages=[],
                        error_message="This PDF is password-protected. Password is required to unlock it.",
                        is_encrypted=True,
                    )

                # Successfully decrypted: create clean unencrypted PDF stream for storage
                try:
                    writer = pypdf.PdfWriter()
                    writer.append(reader)
                    out_stream = io.BytesIO()
                    writer.write(out_stream)
                    unencrypted_content = out_stream.getvalue()
                except Exception as writer_exc:
                    logger.warning("Could not export unencrypted PDF stream: %s", writer_exc)

            pages: list[ExtractedPage] = []
            total_text_length = 0

            for idx, page in enumerate(reader.pages, start=1):
                try:
                    page_text = page.extract_text() or ""
                except Exception:
                    page_text = ""
                cleaned = page_text.strip()
                total_text_length += len(cleaned)
                pages.append(ExtractedPage(page_number=idx, text=cleaned))

            page_count = len(pages)
            # Scanned or image-only PDF detection
            # If total extracted text is minimal across pages, OCR is required
            if page_count > 0 and total_text_length < 20:
                logger.info("PDF has %d pages but insufficient text (%d chars); flagged as OCR_REQUIRED", page_count, total_text_length)
                return ExtractionResult(
                    success=True,
                    status="ocr_required",
                    page_count=page_count,
                    pages=pages,
                    error_message="Image-only or scanned PDF detected. No extractable text stream.",
                    is_encrypted=False,
                    unencrypted_content=unencrypted_content,
                )

            return ExtractionResult(
                success=True,
                status="ready",
                page_count=page_count,
                pages=pages,
                is_encrypted=False,
                unencrypted_content=unencrypted_content,
            )

    except Exception as exc:
        logger.exception("Failed to extract text from %s file: %s", file_type, exc)
        return ExtractionResult(
            success=False,
            status="failed",
            page_count=0,
            pages=[],
            error_message=f"Corrupted or unreadable {file_type.upper()} file.",
        )

    return ExtractionResult(
        success=False,
        status="failed",
        page_count=0,
        pages=[],
        error_message=f"Unsupported format '{file_type}'.",
    )

